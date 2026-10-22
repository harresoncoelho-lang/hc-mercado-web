"""Extração local limitada; não executa anexos nem usa serviços de IA."""
import argparse
import csv
import io
import json
import os
from pathlib import Path, PurePosixPath
import re
import shutil
import signal
import subprocess
import tempfile
import time
import zipfile
import xml.etree.ElementTree as ET

MAX_BYTES = 64 * 1024 * 1024
MAX_ENTRADAS = 200
MAX_PAGINAS = 100
MAX_TEXTO = 600000
OFFICE = {'.doc', '.docx', '.rtf', '.odt', '.xls', '.xlsx', '.ods', '.ppt', '.pptx', '.odp'}
IMAGENS = {'.png', '.jpg', '.jpeg', '.tif', '.tiff', '.bmp', '.webp'}

class Leitor:
    def __init__(self, max_bytes=MAX_BYTES, segundos=90):
        self.max_bytes = max_bytes
        self.restante = max_bytes
        self.prazo = time.monotonic() + segundos
        self.entradas = 0
        self.paginas = 0
        self.textos = []
        self.pendencias = []
        self.lidos = []

    def comando(self, args):
        restante = self.prazo - time.monotonic()
        if restante <= 0:
            raise ValueError('limite de tempo')
        # Não repassar credenciais do worker aos conversores de documentos públicos.
        env = {k: v for k, v in os.environ.items() if k.upper() in {
            'PATH', 'SYSTEMROOT', 'WINDIR', 'TEMP', 'TMP', 'HOME', 'LANG', 'TESSDATA_PREFIX'}}
        inicio = time.monotonic()
        with tempfile.TemporaryFile() as saida, tempfile.TemporaryFile() as erro:
            processo = subprocess.Popen(args, stdout=saida, stderr=erro, stdin=subprocess.DEVNULL,
                                       env=env, start_new_session=(os.name != 'nt'))
            try:
                while processo.poll() is None:
                    if time.monotonic() >= min(self.prazo, inicio + 40):
                        raise ValueError('limite de tempo do conversor')
                    if os.fstat(saida.fileno()).st_size > self.max_bytes or os.fstat(erro.fileno()).st_size > 1024 * 1024:
                        raise ValueError('saída excede limite')
                    time.sleep(0.05)
                if processo.returncode:
                    raise ValueError('conversor falhou (arquivo danificado, senha ou formato incompatível)')
                saida.seek(0)
                dados = saida.read(self.max_bytes + 1)
                if len(dados) > self.max_bytes:
                    raise ValueError('saída excede limite')
                return dados
            finally:
                if processo.poll() is None:
                    if os.name != 'nt':
                        os.killpg(processo.pid, signal.SIGKILL)
                    else:
                        processo.kill()
                    processo.wait()

    def guardar(self, nome, texto):
        texto = texto.strip()
        if not texto:
            raise ValueError('sem texto recuperável')
        rotulo = f'\n--- {nome} ---\n'
        restante = MAX_TEXTO - sum(len(t) for t in self.textos) - len(self.textos) - len(rotulo)
        self.textos.append(rotulo + texto[:max(0, restante)])
        if len(texto) > restante:
            self.pendencias.append(f'{nome}: limite de texto')
        self.lidos.append(nome)

    def ocr(self, caminho):
        return self.comando(['tesseract', str(caminho), 'stdout', '-l', 'por+eng']).decode('utf-8', 'replace')

    def pdf(self, caminho, nome):
        info = self.comando(['pdfinfo', str(caminho)]).decode('utf-8', 'replace')
        match = re.search(r'^Pages:\s+(\d+)', info, re.M)
        if not match:
            raise ValueError('não foi possível contar páginas')
        total = int(match[1])
        limite = min(total, MAX_PAGINAS - self.paginas)
        if limite < total:
            self.pendencias.append(f'{nome}: limite de páginas ({limite}/{total})')
        imagens = self.comando(['pdfimages', '-list', str(caminho)]).decode('utf-8', 'replace')
        paginas_imagem = set()
        for linha in imagens.splitlines():
            campos = linha.split()
            if len(campos) > 4 and campos[0].isdigit() and campos[3].isdigit() and campos[4].isdigit():
                if int(campos[3]) >= 300 and int(campos[4]) >= 300:
                    paginas_imagem.add(int(campos[0]))
        with tempfile.TemporaryDirectory(prefix='anexo-pdf-') as pasta:
            for pagina in range(1, limite + 1):
                self.paginas += 1
                texto = self.comando(['pdftotext', '-f', str(pagina), '-l', str(pagina), '-layout', str(caminho), '-']).decode('utf-8', 'replace')
                if len(texto.strip()) < 80 or pagina in paginas_imagem:
                    try:
                        imagem = Path(pasta) / 'pagina'
                        self.comando(['pdftoppm', '-f', str(pagina), '-l', str(pagina), '-singlefile', '-scale-to', '2500', '-png', str(caminho), str(imagem)])
                        visual = self.ocr(imagem.with_suffix('.png'))
                        if visual.strip():
                            texto += '\n[Texto reconhecido da imagem]\n' + visual
                        elif pagina in paginas_imagem:
                            self.pendencias.append(f'{nome}, página {pagina}: imagem sem texto reconhecido')
                    except (ValueError, OSError, subprocess.SubprocessError):
                        self.pendencias.append(f'{nome}, página {pagina}: OCR não concluído')
                if texto.strip():
                    self.guardar(f'{nome}, página {pagina}', texto)

    def ler(self, caminho, nome=None, nivel=0):
        nome = nome or caminho.name
        try:
            self.entradas += 1
            if self.entradas > MAX_ENTRADAS or nivel > 4:
                raise ValueError('limite de entradas/profundidade')
            if time.monotonic() >= self.prazo:
                raise ValueError('limite de tempo')
            tamanho = caminho.stat().st_size
            if tamanho > self.restante:
                raise ValueError('limite de bytes expandidos')
            self.restante -= tamanho
            with caminho.open('rb') as f:
                assinatura = f.read(8)
            ext = caminho.suffix.lower()
            if assinatura.startswith(b'\x89PNG'):
                ext = '.png'
            elif assinatura.startswith(b'\xff\xd8\xff'):
                ext = '.jpg'
            elif assinatura[:4] in {b'II*\x00', b'MM\x00*'}:
                ext = '.tiff'
            elif assinatura.startswith(b'\xd0\xcf\x11\xe0'):
                if ext not in OFFICE:
                    dados_ole = caminho.read_bytes()
                    ext = '.xls' if any(n.encode('utf-16-le') in dados_ole for n in ('Workbook', 'Book')) else '.doc'
            elif assinatura.startswith(b'BM'):
                ext = '.bmp'
            elif assinatura.startswith(b'RIFF') and caminho.read_bytes()[8:12] == b'WEBP':
                ext = '.webp'
            elif assinatura.startswith(b'{\\rtf'):
                ext = '.rtf'
            if assinatura.startswith(b'%PDF-'):
                self.pdf(caminho, nome)
            elif zipfile.is_zipfile(caminho):
                with zipfile.ZipFile(caminho) as z:
                    if len(z.infolist()) > MAX_ENTRADAS or sum(i.file_size for i in z.infolist()) > self.restante:
                        raise ValueError('limite de expansão do ZIP/Office')
                    nomes = z.namelist()
                    if 'word/document.xml' in nomes:
                        item = z.getinfo('word/document.xml')
                        if item.file_size > self.restante:
                            raise ValueError('limite de bytes do Word')
                        self.restante -= item.file_size
                        raiz = ET.fromstring(z.read(item))
                        ns = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
                        if next(raiz.iter(ns + 'altChunk'), None) is not None:
                            self.pendencias.append(f'{nome}: fragmento externo altChunk exige conferência')
                        texto_word = '\n'.join(''.join(t.text or '' for t in p.iter(ns + 't')) for p in raiz.iter(ns + 'p'))
                        if texto_word.strip():
                            self.guardar(nome, texto_word)
                        for parte in z.infolist():
                            if re.match(r'word/(header|footer|footnotes|endnotes).*\.xml$', parte.filename):
                                if parte.file_size > self.restante:
                                    self.pendencias.append(f'{nome}/{parte.filename}: limite de bytes')
                                    continue
                                self.restante -= parte.file_size
                                raiz_parte = ET.fromstring(z.read(parte))
                                texto_parte = '\n'.join(''.join(t.text or '' for t in p.iter(ns + 't')) for p in raiz_parte.iter(ns + 'p'))
                                if texto_parte.strip():
                                    self.guardar(nome + '/' + parte.filename, texto_parte)
                            elif parte.filename.startswith(('word/media/', 'word/embeddings/')):
                                self.incorporado(z, parte, nome, nivel)
                    elif 'xl/workbook.xml' in nomes:
                        self.planilha(caminho, nome)
                    elif 'mimetype' in nomes and z.read('mimetype')[:100] == b'application/vnd.oasis.opendocument.spreadsheet':
                        with tempfile.TemporaryDirectory(prefix='anexo-ods-') as pasta:
                            ods = Path(pasta) / 'planilha.ods'
                            shutil.copyfile(caminho, ods)
                            self.office(ods, nome)
                    elif '[Content_Types].xml' in nomes or 'mimetype' in nomes:
                        self.office(caminho, nome)
                    else:
                        for item in z.infolist():
                            if item.is_dir():
                                continue
                            if self.entradas >= MAX_ENTRADAS:
                                self.pendencias.append(f'{nome}: limite de entradas')
                                break
                            seguro = PurePosixPath(item.filename.replace('\\', '/'))
                            if seguro.is_absolute() or ':' in item.filename or '..' in seguro.parts or (item.external_attr >> 16) & 0o170000 == 0o120000 or item.file_size > self.restante or item.flag_bits & 1:
                                self.pendencias.append(f'{nome}/{item.filename}: caminho inseguro, senha ou limite de bytes')
                                continue
                            try:
                                with tempfile.TemporaryDirectory(prefix='anexo-zip-') as pasta:
                                    filho = Path(pasta) / ('arquivo' + seguro.suffix)
                                    with z.open(item) as origem, filho.open('wb') as destino:
                                        escrito = 0
                                        while bloco := origem.read(65536):
                                            escrito += len(bloco)
                                            if escrito > self.restante:
                                                raise ValueError('limite de bytes expandidos')
                                            destino.write(bloco)
                                    self.ler(filho, nome + '/' + item.filename, nivel + 1)
                            except (ValueError, OSError, RuntimeError, zipfile.BadZipFile):
                                self.pendencias.append(f'{nome}/{item.filename}: entrada não recuperada')
            elif assinatura.startswith(b'Rar!') or assinatura.startswith(b'7z\xbc\xaf\x27\x1c'):
                self.compactado(caminho, nome, nivel)
            elif ext in OFFICE:
                self.office(caminho, nome, ext)
            elif ext in IMAGENS:
                self.guardar(nome, self.ocr(caminho))
            elif ext in {'.txt', '.csv', '.tsv'}:
                dados = caminho.read_bytes()
                texto = dados.decode('utf-16') if dados.startswith((b'\xff\xfe', b'\xfe\xff')) else self.texto_simples(dados)
                if texto.strip() and ext in {'.csv', '.tsv'}:
                    separador = '\t' if ext == '.tsv' else (';' if ';' in texto.splitlines()[0] else ',')
                    texto = '\n'.join(' | '.join(linha) for linha in csv.reader(io.StringIO(texto), delimiter=separador))
                self.guardar(nome, texto)
            else:
                raise ValueError('formato não suportado')
        except (ValueError, OSError, RuntimeError, subprocess.SubprocessError, zipfile.BadZipFile, ET.ParseError, UnicodeError) as erro:
            self.pendencias.append(f'{nome}: {type(erro).__name__}: {erro}')

    def incorporado(self, pacote, entrada, nome, nivel=0):
        if entrada.file_size > self.restante:
            self.pendencias.append(f'{nome}/{entrada.filename}: limite de bytes incorporados')
            return
        with tempfile.TemporaryDirectory(prefix='anexo-objeto-') as pasta:
            caminho = Path(pasta) / ('objeto' + PurePosixPath(entrada.filename).suffix)
            caminho.write_bytes(pacote.read(entrada))
            self.ler(caminho, nome + '/' + entrada.filename, nivel + 1)

    @staticmethod
    def texto_simples(dados):
        try:
            return dados.decode('utf-8-sig')
        except UnicodeDecodeError:
            return dados.decode('cp1252')

    def planilha(self, caminho, nome):
        from openpyxl import load_workbook
        # Controlar expansão antes da biblioteca abrir o contêiner OOXML.
        with zipfile.ZipFile(caminho) as z:
            tamanho = sum(i.file_size for i in z.infolist())
            if tamanho > self.restante or len(z.infolist()) > MAX_ENTRADAS:
                raise ValueError('limite de expansão da planilha')
            self.restante -= tamanho
            for item in z.infolist():
                if item.filename.startswith(('xl/media/', 'xl/embeddings/')):
                    self.incorporado(z, item, nome)
                elif re.match(r'xl/drawings/[^/]+\.xml$', item.filename):
                    desenho = ET.fromstring(z.read(item))
                    texto_desenho = '\n'.join(t.text or '' for t in desenho.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}t'))
                    if texto_desenho.strip():
                        self.guardar(nome + '/' + item.filename, texto_desenho)
        livro = load_workbook(io.BytesIO(caminho.read_bytes()), read_only=True, data_only=True, keep_links=False)
        formulas = load_workbook(io.BytesIO(caminho.read_bytes()), read_only=True, data_only=False, keep_links=False)
        celulas = 0
        try:
            for folha, fonte in zip(livro.worksheets, formulas.worksheets):
                folha.reset_dimensions()
                fonte.reset_dimensions()
                linhas = []
                for linha, originais in zip(folha.iter_rows(), fonte.iter_rows()):
                    if time.monotonic() >= self.prazo:
                        raise ValueError('limite de tempo da planilha')
                    celulas += len(linha)
                    if celulas > 200000:
                        raise ValueError('limite de células da planilha')
                    valores = []
                    for celula, original in zip(linha, originais):
                        if original.data_type == 'f' and celula.value is None:
                            self.pendencias.append(f'{nome}/{folha.title}/{celula.coordinate}: fórmula sem resultado armazenado')
                        if celula.value is not None:
                            valores.append(f'{celula.coordinate}={celula.value}')
                    if valores:
                        linhas.append(' | '.join(valores))
                if linhas:
                    self.guardar(nome + '/aba ' + folha.title, '\n'.join(linhas))

        finally:
            livro.close()
            formulas.close()

    def office(self, caminho, nome, extensao=None):
        programa = shutil.which('libreoffice') or shutil.which('soffice')
        if not programa:
            raise ValueError('conversor Office indisponível')
        with tempfile.TemporaryDirectory(prefix='anexo-office-') as pasta:
            perfil_pasta = Path(pasta) / 'perfil'
            (perfil_pasta / 'user').mkdir(parents=True)
            (perfil_pasta / 'user' / 'registrymodifications.xcu').write_text(
                '<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry">'
                '<item oor:path="/org.openoffice.Office.Common/Security/Scripting">'
                '<prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop></item>'
                '</oor:items>', encoding='utf-8')
            perfil = perfil_pasta.as_uri()
            formato = 'xlsx' if (extensao or caminho.suffix.lower()) in {'.xls', '.ods'} else 'pdf'
            self.comando([programa, '-env:UserInstallation=' + perfil, '--headless', '--convert-to', formato, '--outdir', pasta, str(caminho)])
            convertido = Path(pasta) / (caminho.stem + '.' + formato)
            if not convertido.is_file() or convertido.stat().st_size > self.restante:
                raise ValueError('conversão Office indisponível ou excede limite')
            if formato == 'xlsx':
                self.planilha(convertido, nome)
            else:
                self.pdf(convertido, nome)

    def compactado(self, caminho, nome, nivel):
        programa = shutil.which('7zz') or shutil.which('7z')
        if not programa:
            raise ValueError('conversor RAR/7z indisponível')
        listagem = self.comando([programa, 'l', '-slt', '-p-', str(caminho)]).decode('utf-8', 'replace').replace('\r\n', '\n')
        blocos = listagem.split('----------\n')
        if len(blocos) < 2:
            raise ValueError('listagem de arquivo indisponível')
        for bloco in blocos[-1].split('\n\n'):
            props = dict(l.split(' = ', 1) for l in bloco.splitlines() if ' = ' in l)
            entrada = props.get('Path')
            if not entrada or props.get('Folder') == '+':
                continue
            if self.entradas >= MAX_ENTRADAS:
                self.pendencias.append(f'{nome}: limite de entradas')
                break
            seguro = PurePosixPath(entrada.replace('\\', '/'))
            tamanho = int(props.get('Size', self.max_bytes + 1))
            if seguro.is_absolute() or ':' in entrada or '..' in seguro.parts or props.get('Symbolic Link') or props.get('Hard Link') or tamanho > self.restante or props.get('Encrypted') == '+':
                self.pendencias.append(f'{nome}/{entrada}: caminho inseguro, senha ou limite de bytes')
                continue
            with tempfile.TemporaryDirectory(prefix='anexo-7z-') as pasta:
                filho = Path(pasta) / ('arquivo' + seguro.suffix)
                filho.write_bytes(self.comando([programa, 'x', '-so', '-spd', '-p-', str(caminho), entrada]))
                self.ler(filho, nome + '/' + entrada, nivel + 1)


def extrair(caminho, max_bytes=MAX_BYTES, segundos=90):
    leitor = Leitor(max_bytes, segundos)
    leitor.ler(Path(caminho))
    return {'texto': '\n'.join(leitor.textos)[:MAX_TEXTO], 'pendencias': leitor.pendencias,
            'lidos': leitor.lidos, 'paginas': leitor.paginas}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('arquivo')
    args = parser.parse_args()
    print(json.dumps(extrair(args.arquivo), ensure_ascii=True))
