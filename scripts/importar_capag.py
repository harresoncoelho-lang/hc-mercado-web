"""Converte a planilha pública do Tesouro em snapshot municipal, sem recalcular notas."""
import argparse
import hashlib
import json
import math
from pathlib import Path

from openpyxl import load_workbook


def converter(linhas):
    registros = []
    for row in linhas:
        ibge, nome, uf, nota, ind1, nota1, ind2, nota2, ind3, nota3, _, observacao, origem = row[:13]
        if not isinstance(ibge, (int, float)) or not nome:
            continue
        if len(str(int(ibge))) != 7 or not isinstance(uf, str) or len(uf) != 2:
            raise ValueError("Identificador municipal inválido na planilha")
        indicadores = []
        for titulo, valor, classificacao in [("Endividamento", ind1, nota1), ("Poupança corrente", ind2, nota2), ("Liquidez relativa", ind3, nota3)]:
            numero = valor if isinstance(valor, (int, float)) and math.isfinite(valor) else None
            indicadores.append({"nome": titulo, "valor": numero, "nota": str(classificacao or "N.D.")})
        registros.append({"ibge": str(int(ibge)), "municipio": nome, "uf": uf, "nota": str(nota or "N.D."), "indicadores": indicadores, "observacaoFonte": observacao, "origemNota": origem})
    if not registros or len({r["ibge"] for r in registros}) != len(registros):
        raise ValueError("Base vazia ou com municípios duplicados")
    return registros


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("arquivo", type=Path)
    p.add_argument("--posicao", required=True)
    p.add_argument("--coletado-em", required=True)
    p.add_argument("--fonte", required=True)
    p.add_argument("--saida", type=Path, default=Path("data/capag-municipios.json"))
    args = p.parse_args()
    if not args.fonte.startswith("https://www.tesourotransparente.gov.br/"):
        raise ValueError("Use a URL oficial do Tesouro")
    w = load_workbook(args.arquivo, read_only=True, data_only=True)
    s = w["Prévia da CAPAG"]
    cabecalho = next(s.iter_rows(min_row=3, max_row=3, values_only=True))
    if cabecalho[:4] != ("Código Município Completo", "Nome_Município", "UF", "CAPAG"):
        raise ValueError("Formato da planilha mudou; revise o importador")
    dados = {"posicao": args.posicao, "coletadoEm": args.coletado_em, "fonte": args.fonte, "sha256": hashlib.sha256(args.arquivo.read_bytes()).hexdigest(), "municipios": converter(s.iter_rows(min_row=4, values_only=True))}
    args.saida.write_text(json.dumps(dados, ensure_ascii=False, allow_nan=False, separators=(",", ":")), encoding="utf-8")
    print(f"Snapshot salvo: {len(dados['municipios'])} municípios; posição {args.posicao}")


if __name__ == "__main__":
    main()
